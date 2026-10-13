"""Excel-Karrieren einlesen und als JSON ablegen.

  python importer/import_excel.py reference/*.xlsx --out data/

Erkennt Tabellen anhand ihrer Überschriften, nicht anhand fester Zellen, damit auch
ältere oder abweichende Saison-Dateien funktionieren. Unklare Stellen werden pro Datei
in einem Bericht aufgelistet.
"""
from __future__ import annotations

import argparse
import json
import re
import sys
from pathlib import Path

import openpyxl

sys.path.insert(0, str(Path(__file__).resolve().parent))
from json_util import clean  # noqa: E402


def norm(s) -> str:
    return re.sub(r"[^a-z0-9]", "", str(s).lower()) if s is not None else ""


def find_header(ws, required: list[str]):
    """Zeile und Spaltenzuordnung der ersten Zeile, die alle Pflicht-Überschriften enthält."""
    req = [norm(r) for r in required]
    for row in ws.iter_rows(min_row=1, max_row=min(ws.max_row, 60)):
        cols = {norm(c.value): c.column for c in row if c.value is not None}
        if all(r in cols for r in req):
            return row[0].row, cols
    return None, None


def read_table(ws, required, stop_on_blank_col=None):
    hdr, cols = find_header(ws, required)
    if hdr is None:
        return [], None
    rows = []
    for r in range(hdr + 1, ws.max_row + 1):
        first = ws.cell(r, cols[norm(stop_on_blank_col or required[0])]).value
        if first is None or str(first).strip() == "":
            if rows:
                break
            continue
        rows.append({k: ws.cell(r, c).value for k, c in cols.items()})
    return rows, hdr


def find_sheet(wb, *keys):
    for ws in wb.worksheets:
        t = norm(ws.title)
        if any(norm(k) in t for k in keys):
            return ws
    return None


def parse_standings(ws, report):
    rows, hdr = read_table(ws, ["Fahrer", "Team", "Punkte"])
    if hdr is None:
        report.append("Fahrerwertung: Tabelle nicht gefunden")
        return [], []
    race_cols = sorted([k for k in rows[0] if re.fullmatch(r"r\d+", k)], key=lambda x: int(x[1:]))
    # Streckennamen stehen in der Zeile unter der Überschrift
    _, cols = find_header(ws, ["Fahrer", "Team", "Punkte"])
    tracks = [ws.cell(hdr + 1, cols[k]).value for k in race_cols]
    calendar = [{"round": i + 1, "track": t} for i, t in enumerate(tracks)]
    drivers = []
    for row in rows:
        if row.get("fahrer") is None or isinstance(row.get("fahrer"), str) and row["fahrer"].strip() == "":
            continue
        results = []
        for k in race_cols:
            v = row[k]
            results.append(None if v in (None, "", "-") else ("DNF" if str(v).upper() == "DNF" else int(v)))
        drivers.append({"name": str(row["fahrer"]).strip(), "team": row.get("team"), "results": results})
    return calendar, drivers


def parse_quali(ws, report):
    rows, hdr = read_table(ws, ["Rennen", "Strecke", "Quali"])
    if hdr is None:
        report.append("Quali & Notizen: Tabelle nicht gefunden")
        return []
    out = []
    for row in rows:
        mate_key = next((k for k in row if k.endswith("q") and k not in ("quali", "qduell")), None)
        q = row.get("quali")
        out.append({
            "round": int(re.sub(r"\D", "", str(row["rennen"])) or 0),
            "country": row.get("land"),
            "track": row.get("strecke"),
            "quali": None if q in (None, "") else (q if isinstance(q, (int, float)) else str(q)),
            "weather": row.get("wetter"),
            "dnf_reason": row.get("dnfgrund"),
            "own_fault": row.get("eigen"),
            "teammate_quali": row.get(mate_key) if mate_key else None,
            "note": next((v for k, v in row.items() if k.startswith("notiz")), None),
        })
    return out


def parse_objectives(ws, report):
    """Ziele in Stufen; Stufenname steht in einer Zeile über der jeweiligen Tabelle."""
    tiers = []
    current = None
    for r in range(1, ws.max_row + 1):
        a = ws.cell(r, 1).value
        if a is None:
            continue
        txt = str(a)
        m = re.match(r"^\W*\s*(PFLICHT|BONUS|KRÖNUNG|KROENUNG)\s+[—-]\s*(.*)$", txt.strip(), re.I)
        if m and "STATUS" not in txt.upper():
            current = {"tier": m.group(1).upper().replace("KROENUNG", "KRÖNUNG"), "title": m.group(2).strip(),
                       "goals": [], "success_text": None, "fail_text": None}
            tiers.append(current)
            continue
        if current and isinstance(a, (int, float)) and ws.cell(r, 2).value:
            current["goals"].append({"goal": ws.cell(r, 2).value, "threshold": ws.cell(r, 3).value})
        elif current and ("erfüllt:" in txt.lower() or "verfehlung:" in txt.lower()):
            if "verfehlung" in txt.lower():
                current["fail_text"] = txt
            else:
                current["success_text"] = txt
    if not tiers:
        report.append("Anforderungen: keine Zielstufen erkannt")
    return tiers


def parse_sections(ws):
    """Freitextabschnitte: Überschrift (Zeile mit Emoji und Großschrift) und folgende Zeilen."""
    sections, current = {}, None
    for r in range(1, ws.max_row + 1):
        a = ws.cell(r, 1).value
        if a is None:
            continue
        txt = str(a).strip()
        if r <= 2:
            sections.setdefault("intro", []).append(txt)
            continue
        if re.match(r"^\W{1,3}\s{2,}[A-ZÄÖÜ][A-ZÄÖÜ &\-]{4,}", txt):
            current = re.sub(r"^\W+\s*", "", txt).strip()
            sections[current] = []
        elif current:
            c = ws.cell(r, 3).value
            sections[current].append(txt if c is None else f"{txt} | {c}")
    return sections


def parse_relationships(sections):
    key = next((k for k in sections if "BEZIEHUNG" in k), None)
    if not key:
        return []
    out = []
    for line in sections[key]:
        m = re.match(r"^(🟢|🟡|🟠|🔴)\s*(.+?)\s+—\s+(.+?)(?:\s+\|\s+(.*))?$", line)
        if m:
            status = {"🟢": "green", "🟡": "yellow", "🟠": "orange", "🔴": "red"}[m.group(1)]
            out.append({"person": m.group(2), "status": status, "label": m.group(3), "text": m.group(4)})
    return out


def parse_rating(sections):
    key = next((k for k in sections if "FAHRERPROFIL" in k), None)
    if not key:
        return {}
    rating = {}
    for line in sections[key]:
        m = re.match(r"^\W*\s*([A-Z]+)\s*\((\d+)\):\s*(.*)$", line)
        if m:
            rating[m.group(1).lower()] = {"value": int(m.group(2)), "text": m.group(3)}
    return rating


def import_file(path: Path) -> dict:
    report = []
    wb = openpyxl.load_workbook(path, data_only=False)
    season = {"source": path.name, "year": None, "player": None, "team": None}
    m = re.search(r"(\d{4})", path.stem)
    if m:
        season["year"] = int(m.group(1))

    ws = find_sheet(wb, "Fahrerwertung")
    calendar, drivers = parse_standings(ws, report) if ws else ([], [])
    ws = find_sheet(wb, "Quali")
    races = parse_quali(ws, report) if ws else []
    ws = find_sheet(wb, "Anforderungen")
    tiers, sections = (parse_objectives(ws, report), parse_sections(ws)) if ws else ([], {})
    ws = find_sheet(wb, "Konstrukteur")
    teams, _ = read_table(ws, ["Team", "Fahrer 1", "Fahrer 2"]) if ws else ([], None)

    title = next((str(ws.cell(1, 1).value) for ws in wb.worksheets if "ANFORDERUNGEN" in str(ws.cell(1, 1).value).upper()), "")
    mt = re.search(r"—\s*(.+?)\s*—\s*([A-ZÄÖÜ][A-ZÄÖÜ\s\.]+?)\s*(\(\d+\))?\s*—", title)
    if mt:
        season["team"] = re.sub(r"\s*\d{4}\s*$", "", mt.group(1)).strip()
        season["player"] = mt.group(2).strip()
    else:
        up = [d["name"] for d in drivers if d["name"].isupper()]
        season["player"] = up[0] if up else None
        if season["player"]:
            season["team"] = next(d["team"] for d in drivers if d["name"] == season["player"])
        report.append("Titelzeile der Anforderungen nicht geparst, Spieler aus Großschreibung abgeleitet")

    for r in races:
        r["track_sub"] = r["track"]
    return clean({
        "season": season,
        "calendar": calendar,
        "drivers": drivers,
        "teams": [{"name": t.get("team"), "driver1": t.get("fahrer1"), "driver2": t.get("fahrer2")} for t in teams],
        "races": races,
        "objectives": tiers,
        "rating": parse_rating(sections),
        "relationships": parse_relationships(sections),
        "sections": sections,
        "report": report,
    })


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("files", nargs="+")
    ap.add_argument("--out", default="data")
    a = ap.parse_args()
    out = Path(a.out); out.mkdir(parents=True, exist_ok=True)
    for f in a.files:
        p = Path(f)
        try:
            data = import_file(p)
        except Exception as e:  # noqa: BLE001
            print(f"FEHLER {p.name}: {e}")
            continue
        target = out / (p.stem + ".json")
        target.write_text(json.dumps(data, ensure_ascii=False, indent=2), encoding="utf-8")
        s = data["season"]
        print(f"{p.name}: {s.get('year')} {s.get('player')} @ {s.get('team')}  "
              f"{len(data['drivers'])} Fahrer, {len(data['races'])} Rennen, "
              f"{len(data['objectives'])} Zielstufen, {len(data['relationships'])} Beziehungen"
              + (f"  HINWEISE: {data['report']}" if data["report"] else ""))


if __name__ == "__main__":
    main()
